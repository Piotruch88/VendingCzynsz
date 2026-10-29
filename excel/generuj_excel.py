import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1]

F = "Arial"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="DDEBF7")
RES = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

ZL = '#,##0" zł"'
ZL2 = '#,##0.00" zł"'
PCT = '0.0%'
NUM = '#,##0'
NUM1 = '#,##0.0'

Z = "'Założenia'!"


def font(cell, bold=False, color="000000", size=10, italic=False):
    cell.font = Font(name=F, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, start_col=2):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.fill = HEAD
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


wb = Workbook()

# ---------------------------------------------------------------- Założenia
za = wb.active
za.title = "Założenia"
kal = wb.create_sheet("Kalkulator", 0)
tab = wb.create_sheet("Tabela")

za["B1"] = "Założenia modelu – czynsz za automaty vendingowe w szkołach"
font(za["B1"], bold=True, size=14)
za["B2"] = ("Niebieskie wartości to założenia – możesz je zmieniać. Numery [Ź..] odsyłają do arkusza Źródła. "
            "Część wartości to szacunki (brak oficjalnych statystyk) – najlepiej skalibruj je na danych z własnych automatów.")
font(za["B2"], italic=True, size=9)

header(za, 3, ["Parametr", "Min\n(scenariusz ostrożny)", "Max\n(scenariusz optymistyczny)", "Jednostka", "Uzasadnienie / źródło"])

scen = [
    (4, "Udział kupujących dziennie (% osób w szkole)", 0.04, 0.07, "%", PCT,
     "SZACUNEK WŁASNY skalibrowany tak, by obrót automatu w szkole ~500 uczniów wynosił ok. 2–4 tys. zł/mies., "
     "zgodnie z danymi branżowymi o obrotach automatów [Ź5, Ź6]. Asortyment ograniczony rozporządzeniem MZ "
     "(Dz.U. 2026 poz. 197, od 1.09.2026) – brak słodyczy i napojów słodzonych [Ź1, Ź2]."),
    (5, "Średnia wartość transakcji brutto", 5.00, 6.00, "zł", ZL2,
     "Branżowo średni koszyk 5–8 zł [Ź4]; w szkole przyjęto dolną część przedziału (woda, sok, kanapka)."),
    (6, "Czynsz jako % obrotu netto automatu", 0.10, 0.15, "%", PCT,
     "Standardowe lokalizacje: 10–20% obrotu [Ź7]; galerie/dworce: 10–15% [Ź8]. Dla szkoły przyjęto 10–15% "
     "(sprzedaż tylko 10 mies., ograniczony asortyment)."),
]
for r, lab, lo, hi, unit, fmt, note in scen:
    za.cell(row=r, column=2, value=lab)
    for col, v in ((3, lo), (4, hi)):
        c = za.cell(row=r, column=col, value=v)
        c.number_format = fmt
        font(c, color=BLUE)
        c.fill = YELLOW
    za.cell(row=r, column=5, value=unit)
    za.cell(row=r, column=6, value=note)

za["B8"] = "Parametry wspólne"
header(za, 8, ["Parametr wspólny", "Wartość", "", "Jednostka", "Uzasadnienie / źródło"])
common = [
    (9, "Personel szkoły jako % liczby uczniów", 0.10, "%", PCT,
     "GUS 2024/25: 511 tys. nauczycieli na ok. 5,3 mln uczących się (≈ 9,6%) [Ź9]; obsługa szkoły dodatkowo."),
    (10, "Dni nauki w miesiącu", 18, "dni", NUM,
     "Rok 2025/26: 1.09.2025–26.06.2026 = 215 dni roboczych, minus święta, przerwy świąteczne, ferie (10) "
     "i dni dyrektorskie (6–10) ≈ 178–182 dni nauki / 10 mies. ≈ 18 [Ź10, wyliczenie własne]."),
    (11, "Miesiące nauki w roku szkolnym", 10, "mies.", NUM,
     "Wrzesień–czerwiec [Ź10]. W wakacje sprzedaż praktycznie zerowa."),
    (12, "Średnia stawka VAT na asortyment", 0.12, "%", PCT,
     "Woda butelkowana 23%, napoje z ≥20% soku, napoje mleczne i żywność 5% [Ź3]. "
     "Mix ok. 40% wody / 60% reszty ≈ 12%. Czynsz liczony od obrotu netto."),
    (13, "Minimalny czynsz za 1 automat (próg)", 150, "zł/mies.", ZL,
     "Najniższe ceny wywoławcze w przetargach szkolnych: 140 zł/automat (Szczecin 2019) [Ź11]. "
     "Sprawdź zarządzenie organu prowadzącego (gmina/powiat) – może określać minimalne stawki."),
    (14, "Maks. liczba transakcji dziennie na 1 automat", 120, "szt./dzień", NUM,
     "SZACUNEK WŁASNY (pojemność + kolejki na przerwach). Powyżej – sugerowany kolejny automat."),
    (15, "Powierzchnia zajmowana przez 1 automat", 1.0, "m²", '0.00" m²"',
     "W przetargach: 1 m² na automat (Szczecin, Białystok) lub 1,5 m² (Opole, Kraków) [Ź12–Ź16, Ź19]."),
    (16, "Zużycie energii przez 1 automat", 200, "kWh/mies.", NUM,
     "Automat z napojami chłodzonymi: 7–14 kWh/dobę wg BSI [Ź17]; nowoczesne automaty mniej. "
     "Przyjęto ok. 6–7 kWh/dobę. Najlepiej: karta katalogowa konkretnego modelu."),
    (17, "Cena energii elektrycznej netto", 1.00, "zł/kWh", ZL2,
     "Taryfy 2026: energia ok. 0,50 zł/kWh + dystrybucja 0,25–0,35 + opłaty ≈ 0,8–0,9 zł/kWh netto, "
     "z opłatami stałymi ok. 1 zł [Ź18]. Podmień na stawkę z faktury szkoły."),
    (18, "Marża brutto operatora (% obrotu netto)", 0.40, "%", PCT,
     "Koszt towaru 50–70% przychodu [Ź7]; marża na napojach 40–60%, przekąskach 30–50% [Ź6]. Przyjęto 40%."),
    (19, "Próg ostrzeżenia: czynsz / marża operatora", 0.40, "%", PCT,
     "SZACUNEK WŁASNY: jeśli czynsz pochłania >40% marży, operatorzy mogą nie składać ofert."),
    (20, "Zaokrąglenie stawek do", 10, "zł", ZL,
     "Stawki wynikowe zaokrąglane do pełnych 10 zł."),
]
for r, lab, v, unit, fmt, note in common:
    za.cell(row=r, column=2, value=lab)
    c = za.cell(row=r, column=3, value=v)
    c.number_format = fmt
    font(c, color=BLUE)
    c.fill = YELLOW
    za.cell(row=r, column=5, value=unit)
    za.cell(row=r, column=6, value=note)

header(za, 22, ["Typ szkoły", "Współczynnik popytu", "", "", "Uzasadnienie (SZACUNKI WŁASNE – brak danych źródłowych)"])
types = [
    ("Szkoła podstawowa", 0.70, "Młodsze dzieci, mniej kieszonkowego, część rodziców zabrania zakupów."),
    ("Liceum ogólnokształcące", 1.00, "Poziom bazowy modelu."),
    ("Technikum", 1.05, "Dłuższe dni zajęć (więcej godzin w szkole)."),
    ("Szkoła branżowa", 0.95, "Część zajęć praktycznych poza szkołą."),
    ("Zespół szkół", 0.95, "Mieszany profil uczniów."),
    ("Szkoła policealna / dla dorosłych", 1.10, "Dorośli słuchacze, wyższa siła nabywcza."),
]
for i, (name, k, note) in enumerate(types):
    r = 23 + i
    za.cell(row=r, column=2, value=name)
    c = za.cell(row=r, column=3, value=k)
    c.number_format = '0.00'
    font(c, color=BLUE)
    c.fill = YELLOW
    za.cell(row=r, column=6, value=note)
TYPE_NAMES = f"{Z}$B$23:$B$28"
TYPE_K = f"{Z}$C$23:$C$28"

for row in za.iter_rows(min_row=4, max_row=28, min_col=2, max_col=6):
    for c in row:
        if c.font.color is None or c.font.color.rgb in (None, "FF000000"):
            font(c)
        if c.column in (2, 6, 5):
            c.alignment = Alignment(vertical="center", wrap_text=True)
        if c.value is not None or c.column in (3, 4):
            c.border = BOX
for col, w in zip("ABCDEF", (2, 44, 18, 18, 11, 90)):
    za.column_dimensions[col].width = w
za.freeze_panes = "B4"

# ---------------------------------------------------------------- Kalkulator
k = kal
k["B1"] = "Kalkulator czynszu za automaty vendingowe w szkole"
font(k["B1"], bold=True, size=16)
k["B2"] = "Wpisz liczbę uczniów (żółte pole) – model wylicza minimalną i maksymalną stawkę czynszu netto."
font(k["B2"], italic=True, size=9)

header(k, 4, ["DANE WEJŚCIOWE", "Wartość", "", "Wskazówka"])
inputs = [
    (5, "Liczba uczniów w szkole", 500, NUM, "Wpisz liczbę uczniów (np. z SIO lub strony szkoły)."),
    (6, "Typ szkoły", "Liceum ogólnokształcące", None, "Wybierz z listy – wpływa na popyt (domyślnie liceum)."),
    (7, "Liczba automatów w szkole", 1, NUM, "Ile automatów ma stanąć. Czynsz łączny dzielony jest na automaty."),
]
for r, lab, v, fmt, hint in inputs:
    font(k.cell(row=r, column=2, value=lab), bold=True)
    c = k.cell(row=r, column=3, value=v)
    font(c, bold=True, color=BLUE, size=11)
    c.fill = YELLOW
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
    font(k.cell(row=r, column=5, value=hint), italic=True, size=9)
    k.cell(row=r, column=2).border = BOX

dv_n = DataValidation(type="whole", operator="between", formula1="1", formula2="10000",
                      showErrorMessage=True, errorTitle="Błędna wartość",
                      error="Podaj liczbę całkowitą uczniów (1–10 000).")
dv_t = DataValidation(type="list", formula1=f"={TYPE_NAMES}", showErrorMessage=True)
dv_m = DataValidation(type="whole", operator="between", formula1="1", formula2="20",
                      showErrorMessage=True, error="Podaj liczbę automatów (1–20).")
for dv, ref in ((dv_n, "C5"), (dv_t, "C6"), (dv_m, "C7")):
    k.add_data_validation(dv)
    dv.add(ref)

R = f"{Z}$C$20"  # rounding step

# --- calculation block (rows 21-33)
header(k, 20, ["SZCZEGÓŁY OBLICZEŃ", "Scenariusz MIN", "Scenariusz MAX", "Opis"])
calc = [
    (21, "Osoby korzystające (uczniowie + personel)",
     f"=$C$5*(1+{Z}$C$9)", "=C21", NUM, "Uczniowie powiększeni o personel szkoły."),
    (22, "Współczynnik typu szkoły",
     f"=IFERROR(INDEX({TYPE_K},MATCH($C$6,{TYPE_NAMES},0)),1)", "=C22", "0.00", "Z tabeli w arkuszu Założenia."),
    (23, "Udział kupujących dziennie",
     f"={Z}C4", f"={Z}D4", PCT, "Założenie min / max."),
    (24, "Transakcje dziennie – popyt",
     "=C21*C22*C23", "=D21*D22*D23", NUM1, "Ile zakupów dziennie wygeneruje szkoła."),
    (25, "Transakcje dziennie – po limicie wydajności",
     f"=MIN(C24,$C$7*{Z}$C$14)", f"=MIN(D24,$C$7*{Z}$C$14)", NUM1,
     "Popyt ograniczony liczbą automatów."),
    (26, "Średnia wartość transakcji brutto",
     f"={Z}C5", f"={Z}D5", ZL2, "Założenie min / max."),
    (27, "Obrót dzienny brutto", "=C25*C26", "=D25*D26", ZL, ""),
    (28, "Obrót miesięczny brutto", f"=C27*{Z}$C$10", f"=D27*{Z}$C$10", ZL, "× dni nauki w miesiącu."),
    (29, "Obrót miesięczny netto", f"=C28/(1+{Z}$C$12)", f"=D28/(1+{Z}$C$12)", ZL, "Bez VAT."),
    (30, "Czynsz jako % obrotu netto", f"={Z}C6", f"={Z}D6", PCT, "Założenie min / max."),
    (31, "Czynsz wyliczony z obrotu", "=C29*C30", "=D29*D30", ZL, ""),
    (32, "Próg minimalny (× liczba automatów)", f"={Z}$C$13*$C$7", f"={Z}$C$13*$C$7", ZL,
     "Czynsz nie spada poniżej progu."),
    (33, "Czynsz miesięczny łącznie (zaokrąglony)",
     f"=ROUND(MAX(C31,C32)/{R},0)*{R}", f"=ROUND(MAX(D31,D32,C33)/{R},0)*{R}", ZL,
     "Większa z wartości: z obrotu lub próg."),
]
for r, lab, fmin, fmax, fmt, desc in calc:
    font(k.cell(row=r, column=2, value=lab))
    for col, f in ((3, fmin), (4, fmax)):
        c = k.cell(row=r, column=col, value=f)
        c.number_format = fmt
        is_link = f.startswith(f"={Z}") and "*" not in f
        font(c, color=GREEN if is_link else "000000", bold=(r == 33))
        c.border = BOX
    font(k.cell(row=r, column=5, value=desc), italic=True, size=9)
    k.cell(row=r, column=2).border = BOX
for col in (2, 3, 4):
    k.cell(row=33, column=col).fill = SUB
font(k["B33"], bold=True)

# --- results block (rows 9-18)
header(k, 9, ["WYNIK (kwoty netto)", "MINIMALNA", "MAKSYMALNA", "Komentarz"])
res = [
    (10, "Czynsz miesięczny łącznie", "=C33", "=D33", ZL,
     "Stawka za wszystkie automaty w szkole, zł netto / miesiąc nauki."),
    (11, "Stawka wyjściowa do negocjacji / przetargu",
     f"=ROUND(AVERAGE(C10:D10)/{R},0)*{R}", None, ZL,
     "Środek przedziału – rozsądna cena wywoławcza."),
    (12, "Czynsz na 1 automat", "=C10/$C$7", "=D10/$C$7", ZL, ""),
    (13, "Czynsz za 1 m² powierzchni", f"=C12/{Z}$C$15", f"=D12/{Z}$C$15", ZL2,
     "Przydatne, gdy organ prowadzący określa stawki za m²."),
    (14, "Czynsz za rok szkolny", f"=C10*{Z}$C$11", f"=D10*{Z}$C$11", ZL,
     "Przy płatności za miesiące nauki (wrzesień–czerwiec)."),
    (15, "Odpowiednik przy płatności przez 12 miesięcy", "=C14/12", "=D14/12", ZL,
     "Gdy umowa przewiduje równą ratę także w wakacje."),
    (16, "Ryczałt za energię (doliczany osobno)", f"={Z}$C$16*{Z}$C$17*$C$7", None, ZL,
     "Wszystkie automaty, zł netto / mies. – jeśli prąd rozliczany poza czynszem."),
    (17, "Sugerowana liczba automatów", f"=MAX(1,ROUNDUP(D24/{Z}$C$14,0))", None, NUM, None),
    (18, "Udział czynszu w marży brutto operatora",
     f"=IFERROR(C10/(C29*{Z}$C$18),0)", f"=IFERROR(D10/(D29*{Z}$C$18),0)", PCT, None),
]
for r, lab, fmin, fmax, fmt, desc in res:
    font(k.cell(row=r, column=2, value=lab), bold=r in (10, 11))
    k.cell(row=r, column=2).border = BOX
    for col, f in ((3, fmin), (4, fmax)):
        c = k.cell(row=r, column=col, value=f)
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        if f is None:
            continue
        c.number_format = fmt
        font(c, bold=r in (10, 11), size=14 if r == 10 else (12 if r == 11 else 10))
        c.fill = RES
    if desc:
        font(k.cell(row=r, column=5, value=desc), italic=True, size=9)
for r in (11, 16, 17):
    k.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)

k["E17"] = ('=IF($C$7<C17,"Popyt przekracza wydajność – rozważ "&C17&" automat(y) lub ceny wzrosną ponad model",'
            '"Liczba automatów wystarczająca")')
k["E18"] = (f'=IF(MAX(C18:D18)>{Z}$C$19,"UWAGA: czynsz pochłania dużą część marży – ryzyko braku ofert operatorów",'
            '"OK – stawki realne dla operatora")')
for ref in ("E17", "E18"):
    font(k[ref], italic=True, size=9, bold=True)
k.conditional_formatting.add("E17", FormulaRule(formula=['$C$7<$C$17'], font=Font(color="C00000")))
k.conditional_formatting.add("E17", FormulaRule(formula=['$C$7>=$C$17'], font=Font(color="008000")))
k.conditional_formatting.add("E18", FormulaRule(formula=['LEFT($E$18,2)<>"OK"'], font=Font(color="C00000")))
k.conditional_formatting.add("E18", FormulaRule(formula=['LEFT($E$18,2)="OK"'], font=Font(color="008000")))

k["B10"].comment = Comment("Wynik główny modelu. Stawki netto (do czynszu dolicza się 23% VAT).", "Model")

# --- legend
lg = 36
font(k.cell(row=lg, column=2, value="Legenda i uwagi"), bold=True)
legend = [
    ("Żółte pole, niebieska czcionka", "dane do wpisania / założenia, które możesz zmieniać"),
    ("Zielona czcionka", "wartość pobrana z arkusza Założenia"),
    ("Czarna czcionka", "formuły – nie edytuj"),
    ("Kwoty", "netto; do czynszu dolicza się 23% VAT, prąd zwykle jako osobny ryczałt"),
    ("Metoda", "czynsz = obrót netto automatu × % dla szkoły; obrót = osoby × % kupujących × wartość zakupu × dni nauki"),
    ("Kalibracja", "jeśli masz dane z własnych automatów, popraw w Założeniach % kupujących i wartość transakcji"),
    ("Prawo", "sprawdź zarządzenie organu prowadzącego (gmina/powiat) – może narzucać minimalne stawki za m²"),
]
for i, (a, b) in enumerate(legend):
    r = lg + 1 + i
    ca = k.cell(row=r, column=2, value=a)
    font(ca, bold=True, size=9)
    font(k.cell(row=r, column=3, value=b), size=9)
k.cell(row=lg + 1, column=2).fill = YELLOW
k.cell(row=lg + 1, column=2).font = Font(name=F, bold=True, size=9, color=BLUE)
k.cell(row=lg + 2, column=2).font = Font(name=F, bold=True, size=9, color=GREEN)

for col, w in zip("ABCDE", (2, 46, 20, 20, 70)):
    k.column_dimensions[col].width = w
k.row_dimensions[10].height = 24
k.row_dimensions[4].height = 20
k.row_dimensions[9].height = 20
k.row_dimensions[20].height = 20
k.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Tabela
t = tab
t["A1"] = "Tabela stawek czynszu wg liczby uczniów"
font(t["A1"], bold=True, size=14)
t["A2"] = "Typ szkoły (z Kalkulatora):"
font(t["A2"], size=9)
t["C2"] = "=Kalkulator!C6"
font(t["C2"], bold=True, color=GREEN)
t["E2"] = "Współczynnik:"
font(t["E2"], size=9)
t["F2"] = f"=IFERROR(INDEX({TYPE_K},MATCH(C2,{TYPE_NAMES},0)),1)"
t["F2"].number_format = "0.00"
font(t["F2"], bold=True)
t["A3"] = "Liczba automatów = sugerowana dla danej wielkości szkoły. Kwoty netto, zł / miesiąc nauki."
font(t["A3"], italic=True, size=9)

header(t, 5, ["Liczba uczniów", "Sugerowana liczba automatów", "Czynsz MIN\n(łącznie)",
              "Czynsz MAX\n(łącznie)", "Stawka wyjściowa", "MIN na 1 automat", "MAX na 1 automat"], start_col=1)
t.row_dimensions[5].height = 32
students = [50, 100, 150, 200, 250, 300, 400, 500, 600, 700, 800, 900, 1000, 1200, 1500, 2000]
for i, n in enumerate(students):
    r = 6 + i
    users = f"A{r}*(1+{Z}$C$9)*$F$2"
    c = t.cell(row=r, column=1, value=n)
    font(c, color=BLUE, bold=True)
    c.number_format = NUM
    t.cell(row=r, column=2, value=f"=MAX(1,ROUNDUP({users}*{Z}$D$4/{Z}$C$14,0))")
    for col, sc in ((3, "C"), (4, "D")):
        raw = (f"MIN({users}*{Z}${sc}$4,B{r}*{Z}$C$14)*{Z}${sc}$5*{Z}$C$10/(1+{Z}$C$12)*{Z}${sc}$6")
        t.cell(row=r, column=col, value=f"=ROUND(MAX({raw},{Z}$C$13*B{r})/{R},0)*{R}")
    t.cell(row=r, column=5, value=f"=ROUND(AVERAGE(C{r}:D{r})/{R},0)*{R}")
    t.cell(row=r, column=6, value=f"=C{r}/B{r}")
    t.cell(row=r, column=7, value=f"=D{r}/B{r}")
    for col in range(1, 8):
        cell = t.cell(row=r, column=col)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
        if col >= 2:
            font(cell, bold=(col == 5))
            cell.number_format = NUM if col == 2 else ZL
        if i % 2:
            cell.fill = PatternFill("solid", fgColor="F2F2F2")
last = 6 + len(students) - 1
t.cell(row=last + 2, column=1, value="Liczby uczniów w kolumnie A możesz zmienić (niebieskie) – stawki przeliczą się same.")
font(t.cell(row=last + 2, column=1), italic=True, size=9)
for col, w in zip("ABCDEFG", (16, 16, 16, 16, 16, 16, 16)):
    t.column_dimensions[col].width = w
t.freeze_panes = "A6"

ch = LineChart()
ch.title = "Czynsz miesięczny (netto) a liczba uczniów"
ch.y_axis.title = "zł / miesiąc"
ch.x_axis.title = "Liczba uczniów"
ch.height = 9
ch.width = 18
data = Reference(t, min_col=3, max_col=5, min_row=5, max_row=last)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(t, min_col=1, min_row=6, max_row=last))
ch.x_axis.delete = False
ch.y_axis.delete = False
t.add_chart(ch, "I5")

# ---------------------------------------------------------------- Źródła
src = wb.create_sheet("Źródła")
src["A1"] = "Źródła danych i założeń modelu"
font(src["A1"], bold=True, size=14)
src["A2"] = ("Stan na 29.10.2026. Numery [Ź..] są przywoływane w arkuszu Założenia. Strony firm vendingowych to "
             "materiały branżowe (mogą być optymistyczne) – traktuj je jako orientację, nie dane oficjalne.")
font(src["A2"], italic=True, size=9)

COLS = ["Nr", "Temat", "Co mówi źródło (kluczowe liczby)", "Źródło", "Data", "Link", "Jak użyto w modelu / uwagi"]
header(src, 4, COLS, start_col=1)
sources = [
    ("Ź1", "Asortyment w szkołach – prawo",
     "Nowe rozporządzenie MZ z 16.02.2026 w sprawie grup środków spożywczych przeznaczonych do sprzedaży dzieciom "
     "i młodzieży w jednostkach systemu oświaty; wchodzi w życie 1.09.2026, uchyla rozporządzenie z 26.07.2016.",
     "Dz.U. 2026 poz. 197 (infor.pl – tekst aktu)", "20.02.2026",
     "https://www.infor.pl/akt-prawny/DZU.2026.051.0000197,rozporzadzenie-ministra-zdrowia-w-sprawie-grup-srodkow-spozywczych-przeznaczonych-do-sprzedazy-dzieciom-i-mlodziezy-w-jednostkach-systemu-oswiaty-oraz-wymagan-jakie-musza-spelniac-srodki-spozywcze-sto.html",
     "Uzasadnia niższy udział kupujących niż w biurach/galeriach."),
    ("Ź2", "Asortyment w automatach szkolnych",
     "Zamknięta lista grup produktów obowiązuje także w automatach vendingowych. Dozwolone m.in.: woda, soki bez "
     "dodatku cukru, kanapki, owoce, orzechy, nabiał, gorzka czekolada. Brak zwykłych słodyczy i napojów słodzonych.",
     "PSSE w Turku (gov.pl)", "02.03.2026",
     "https://www.gov.pl/web/psse-turek/nowe-rozporzadzenie-w-sprawie-zywienia-dzieci-i-mlodziezy-w-szkolach-i-przedszkolach",
     "j.w. – ograniczony asortyment = niższa sprzedaż i niższy średni koszyk."),
    ("Ź3", "Stawki VAT",
     "Woda butelkowana – 23%. Napoje z ≥20% soku, napoje mleczne i roślinne – 5%.",
     "Infor.pl / PAP (Księgowość)", "14.09.2026",
     "https://ksiegowosc.infor.pl/wiadomosci/7511182,23-vat-na-wode-butelkowana-5-vat-na-napoje-gazowane-i-dosladzane-z-dodatkiem-soku-dlaczego-polski-system-podatkowy-dyskryminuje-wode.html",
     "Średni VAT asortymentu przyjęty 12% (mix wody 23% i reszty 5%)."),
    ("Ź4", "Transakcje i średni koszyk",
     "40–80 transakcji dziennie (zależnie od lokalizacji), średni koszyk 5–8 zł, koszt towaru ok. 50% przychodu.",
     "VendHub – „Ile można zarobić na automatach vendingowych”", "19.03.2025",
     "https://www.vendhub.pl/32-ile-mozna-zarobic-na-automatach-vendingowych-przykladowe-kalkulacje/",
     "Średni koszyk 5–6 zł (dolna część przedziału – szkoła). Dane dotyczą biur/galerii/uczelni, nie szkół."),
    ("Ź5", "Zarobki automatu w szkole",
     "Szkoła: ok. 800 zł/mies. (napoje), ok. 1 000 zł/mies. (przekąski); uczelnia: 1 500–1 700 zł/mies. "
     "Artykuł nie precyzuje, czy to zysk czy obrót.",
     "automatyvending.waw.pl – „Ile można zarobić na automatach vendingowych”", "brak daty",
     "https://automatyvending.waw.pl/ile-mozna-zarobic-na-automatach-vendingowych/",
     "Kalibracja udziału kupujących (4–7%)."),
    ("Ź6", "Obroty i marże",
     "Obrót: słabszy punkt 1 000–1 500 zł/mies., ruchliwa lokalizacja 3 000–6 000 zł/mies. "
     "Marża brutto: napoje butelkowane 40–60%, przekąski pakowane 30–50%.",
     "zwiadowca.pl – „Czy biznes vendingowy się opłaca”", "06.01.2026",
     "https://zwiadowca.pl/informacje/czy-sie-oplaca-zarobki-marze",
     "Kalibracja obrotu; marża brutto operatora 40%."),
    ("Ź7", "Czynsz / prowizja dla lokalizacji",
     "Dzierżawa w standardowych lokalizacjach 10–20% przychodu, w premium 30–50%; kwotowo często od 200–400 zł/mies. "
     "Koszt towaru 50–70% przychodu. Przychód: dobra lokalizacja 800–2 000 zł/mies., uczelnie 3 000–5 000 zł.",
     "nssd.pl – „Automaty vendingowe czy się opłaca”", "10.05.2026",
     "https://nssd.pl/automaty-vendingowe-czy-sie-oplaca/",
     "Czynsz 10–15% obrotu netto. Uwaga: składowe kosztów w artykule nie sumują się spójnie."),
    ("Ź8", "Czynsz / prowizja dla lokalizacji",
     "Galerie i dworce: zwykle 10–15% obrotu lub 500–1 000 zł/mies.; mniejsze placówki publiczne: 100–300 zł/m².",
     "E-lite – „Ile zysku przynosi automat vendingowy”", "29.01.2026 (akt. 12.02.2026)",
     "https://e-lite.com.pl/ile-zysku-przynosi-automat-vendingowy-kalkulator-rentownosci/",
     "Potwierdza przedział 10–15% i rząd wielkości kwot."),
    ("Ź9", "Personel szkół",
     "Rok szkolny 2024/25: ok. 5,3 mln uczących się, 511 tys. nauczycieli.",
     "GUS – „Polska szkoła w liczbach” (Oświata i wychowanie 2024/2025)", "2025",
     "https://stat.gov.pl/dla-mediow/informacje-prasowe/polska-szkola-w-liczbach-jak-wyglada-edukacja-w-roku-szkolnym-20242025,36,1.html",
     "Personel = 10% liczby uczniów."),
    ("Ź10", "Kalendarz roku szkolnego",
     "Rok 2025/26: 1.09.2025–26.06.2026; dni dyrektorskie: 6 (SP) do 10 (LO, technika).",
     "Forsal.pl", "01.07.2025",
     "https://forsal.pl/kraj/aktualnosci/artykuly/9830886,wyjatkowo-niekorzystny-uklad-roku-szkolnego-20252026-nie-beda-zadowo.html",
     "Wyliczenie własne: ok. 178–182 dni nauki → 18 dni/mies. przez 10 mies."),
    ("Ź11", "Przetarg – szkoła",
     "ZS Łączności w Szczecinie, 2 automaty: cena wywoławcza 140 zł netto/automat/mies., najwyższa oferta 230 zł netto.",
     "BIP Technikum Łączności i Multimediów Cyfrowych w Szczecinie", "09.2019",
     "https://bip.tlimc.szczecin.pl/index.php?id=30",
     "Próg minimalny 150 zł. Dane historyczne (2019)."),
    ("Ź12", "Przetarg – szkoła",
     "ZS nr 4 w Szczecinie (ogł. 2/2025): powierzchnia 100×100 cm pod automaty z napojami i przekąskami, "
     "minimalna cena 400 zł netto/mies. + VAT.",
     "Ogłoszenie ZS nr 4 Szczecin", "02.2025",
     "https://www.zs4.szczecin.pl/wp-content/uploads/2025/02/Przetarg-na-automaty.pdf",
     "Punkt odniesienia dla stawek; powierzchnia 1 m²."),
    ("Ź13", "Przetarg – szkoła",
     "ZS nr 4 w Szczecinie (ogł. 3/2025): automat z gorącymi napojami (bez kofeiny), minimalna cena 400 zł netto/mies. + VAT.",
     "Ogłoszenie ZS nr 4 Szczecin", "03.2025",
     "https://www.zs4.szczecin.pl/wp-content/uploads/2025/02/Przetarg-na-automaty-.pdf",
     "Punkt odniesienia dla stawek."),
    ("Ź14", "Przetarg – szkoła",
     "ZS Gastronomicznych w Białymstoku: 2 automaty, 2 m², cena wywoławcza 800 zł netto/mies. (400 zł/automat). "
     "Status: rozstrzygnięty (wylicytowana kwota niepodana na stronie).",
     "BIP UM Białystok", "13.06.2025",
     "https://www.bip.bialystok.pl/postepowania/przetargi_na_nieruchomosci/zespol-szkol-gastronomicznych-w-bialymstoku-ul-knyszynska-12-1.html",
     "Punkt odniesienia; 1 m² na automat."),
    ("Ź15", "Przetarg – szkoła",
     "ZS Budowlano-Geodezyjnych w Białymstoku: 2 automaty, 2 m², cena wywoławcza 1 350 zł netto/mies. (675 zł/automat). "
     "Status: NIEROZSTRZYGNIĘTY (brak rozstrzygnięcia).",
     "BIP UM Białystok", "28.08.2026",
     "https://www.bip.bialystok.pl/postepowania/przetargi_na_nieruchomosci/zespol-szkol-budowlano-geodezyjnych-im-stefana-wladyslawa-bryly-w-bialymstoku-ul-slonimska-471-1.html",
     "Sygnał, że ~675 zł/automat to dla operatorów za dużo w tej szkole – górna granica rynku."),
    ("Ź16", "Przetarg – szpital (porównawczo)",
     "Uniwersytecki Szpital Kliniczny w Opolu: 4 automaty po 1,5 m², min. 250 zł netto/automat/mies.; "
     "energia i woda płacone przez najemcę osobno.",
     "BIP USK Opole", "2024",
     "https://bip.usk.opole.pl/download/attachment/6483/1-ogloszenie-najem-powierzchni-pod-automaty-vendingowe-2024-aktualizacja-czynszu-1.pdf",
     "Media rozliczane poza czynszem (jak ryczałt za energię w modelu)."),
    ("Ź17", "Zużycie energii",
     "Automat z napojami chłodzonymi zużywa 7–14 kWh/dobę; norma pomiaru BS EN IEC 63252:2020.",
     "BSI Knowledge – BS EN IEC 63252", "2020",
     "https://knowledge.bsigroup.com/products/energy-consumption-of-vending-machines-1",
     "Przyjęto 200 kWh/mies. (nowsze automaty zużywają mniej)."),
    ("Ź18", "Cena energii 2026",
     "Średnia cena energii w taryfach 2026: 495,16 zł/MWh; dystrybucja zmienna 0,25–0,35 zł/kWh + opłaty "
     "(OZE, kogeneracyjna, akcyza, jakościowa, mocowa).",
     "Murator Dom – „Ile zapłacisz za prąd w 2026 r.”", "04.01.2026",
     "https://muratordom.pl/prawo/finanse/ile-zaplacisz-za-prad-w-2026-r-aa-1t2X-FBQU-t5cM.html",
     "≈ 0,8–0,9 zł/kWh netto + opłaty stałe → przyjęto 1,00 zł/kWh. Taryfy dla gospodarstw (G11) – szkoła ma inną."),
    ("Ź19", "Przetarg – szkoła",
     "XVII LO (ZSO nr 7) w Krakowie: 1 automat; oferta w zł brutto za m²; gdy brak pomiaru przyjmuje się 1,5 m².",
     "Strona XVII LO w Krakowie", "20.01.2025",
     "https://liceum.zso7.pl/najem-powierzchni/",
     "Powierzchnia automatu; brak stawki minimalnej w ogłoszeniu."),
    ("Ź20", "Przetarg – uczelnia (porównawczo)",
     "Uniwersytet Medyczny w Poznaniu: 1,5 m² pod 1 automat, min. 400 zł netto/mies., w lipcu–wrześniu 50% stawki.",
     "Platforma Zakupowa (opis postępowania z wyników wyszukiwania)", "b.d.",
     "https://platformazakupowa.pl/transakcja/1095271",
     "NIEZWERYFIKOWANE bezpośrednio (strona wymaga logowania). Pomysł obniżki w wakacje."),
    ("Ź21", "Przetarg – szkoła (historyczny)",
     "ZSP w Świdwinie: 3 m², do 2 automatów, min. 70 zł netto/m², energia wliczona w czynsz.",
     "BIP ZS w Świdwinie", "21.06.2012",
     "https://bip.zsswidwin.pl/zamowienia/pokaz/12",
     "Tylko historycznie – pokazuje wariant z mediami wliczonymi w czynsz."),
]
LINK_FONT = Font(name=F, size=9, color="0563C1", underline="single")
for i, row in enumerate(sources):
    r = 5 + i
    for j, v in enumerate(row):
        c = src.cell(row=r, column=1 + j, value=v)
        font(c, size=9, bold=(j == 0))
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX
        if j == 5:
            c.hyperlink = v
            c.font = LINK_FONT
    if i % 2:
        for j in range(7):
            src.cell(row=r, column=1 + j).fill = PatternFill("solid", fgColor="F2F2F2")
end = 5 + len(sources)

r = end + 1
font(src.cell(row=r, column=1, value="Porównanie modelu z przetargami"), bold=True, size=11)
src.cell(row=r + 1, column=1,
         value='="Model (typ: "&Kalkulator!C6&", 500 uczniów, 1 automat): "&TEXT(Tabela!C13,"0")&"–"&TEXT(Tabela!D13,"0")'
               '&" zł netto/mies., stawka wyjściowa "&TEXT(Tabela!E13,"0")&" zł."')
font(src.cell(row=r + 1, column=1), bold=True, size=9)
src.cell(row=r + 2, column=1,
         value="Przetargi szkolne: ceny wywoławcze 140–675 zł/automat (Ź11–Ź15); wylicytowano 230 zł (2019); "
               "przetarg z ceną 675 zł/automat nie znalazł najemcy (Ź15). Liczby uczniów tych szkół nie są znane, "
               "więc to kontrola rzędu wielkości, a nie dokładna kalibracja.")
font(src.cell(row=r + 2, column=1), size=9)

r += 4
font(src.cell(row=r, column=1, value="Założenia bez źródła (szacunki własne – do weryfikacji)"), bold=True, size=11)
own = [
    "Udział kupujących dziennie 4–7% – brak publicznych danych; dobrany tak, by obrót był zgodny z Ź5–Ź7.",
    "Wydajność 120 transakcji dziennie na automat – szacunek (pojemność, kolejki na przerwach).",
    "Współczynniki typu szkoły (SP 0,70 … policealna 1,10) – szacunek.",
    "Próg ostrzeżenia 40% marży operatora – szacunek.",
    "Najlepsza kalibracja: dane sprzedażowe z własnych automatów lub informacje o wylicytowanych stawkach "
    "w przetargach w Twoim regionie (BIP szkół / organu prowadzącego).",
]
for i, t_ in enumerate(own):
    font(src.cell(row=r + 1 + i, column=1, value="• " + t_), size=9)

for col, w in zip("ABCDEFG", (6, 22, 60, 30, 13, 40, 45)):
    src.column_dimensions[col].width = w
src.freeze_panes = "A5"

wb.active = 0
wb.save(OUT)
print("saved", OUT)
